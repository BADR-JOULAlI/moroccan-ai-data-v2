"""Collecte autonome de pages provenant de sites a contenu principalement marocain.

Commande recommandee :

    uv run python collect_moroccan_websites.py --max-links 1000

Le script ne collecte pas les plateformes sociales ou les depots de datasets.
Il produit directement, a la racine du projet :

    moroccan_websites.xlsx
    moroccan_websites.csv
    moroccan_websites.jsonl
"""

from __future__ import annotations

import argparse
import asyncio
import csv
import json
import logging
import re
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import parse_qs, parse_qsl, quote_plus, unquote, urlencode, urljoin, urlparse, urlunparse

import httpx
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


LOGGER = logging.getLogger("moroccan_websites")
ROOT = Path(__file__).resolve().parent

OUTPUT_CSV = ROOT / "moroccan_websites.csv"
OUTPUT_JSONL = ROOT / "moroccan_websites.jsonl"
OUTPUT_XLSX = ROOT / "moroccan_websites.xlsx"

USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) "
    "Chrome/125.0.0.0 Safari/537.36"
)

# Sources marocaines connues servant de points de depart. Le moteur de recherche
# en decouvre d'autres, puis le validateur controle leur contenu avant export.
SEED_SITES = [
    "https://www.maroc.ma/",
    "https://www.mapexpress.ma/",
    "https://www.minculture.gov.ma/",
    "https://www.archivesdumaroc.ma/",
    "https://www.bnrm.ma/",
    "https://fnm.ma/",
    "https://www.ircam.ma/",
    "https://www.habous.gov.ma/",
    "https://www.hcp.ma/",
    "https://www.cg.gov.ma/",
    "https://www.parlement.ma/",
    "https://www.data.gov.ma/",
    "https://www.onmt.ma/",
    "https://www.visitmorocco.com/",
    "https://www.hespress.com/",
    "https://www.le360.ma/",
    "https://www.2m.ma/",
    "https://snrtnews.com/",
    "https://medias24.com/",
    "https://telquel.ma/",
    "https://lematin.ma/",
    "https://aujourdhui.ma/",
    "https://www.yabiladi.com/",
    "https://www.bladi.net/",
    "https://alyaoum24.com/",
    "https://al3omk.com/",
    "https://www.goud.ma/",
    "https://barlamane.com/",
    "https://lakome2.com/",
    "https://achkayen.com/",
    "https://chouftv.ma/",
    "https://kech24.com/",
    "https://rue20.com/",
    "https://anfaspress.com/",
    "https://ar.hibapress.com/",
    "https://bayanealyaoume.press.ma/",
    "https://www.libe.ma/",
    "https://moroccoworldnews.com/",
    "https://en.yabiladi.com/",
    "https://zamane.ma/",
    "https://www.challenge.ma/",
    "https://lobservateur.info/",
]

TRUSTED_SEED_DOMAINS = {
    urlparse(url).netloc.lower().removeprefix("www.") for url in SEED_SITES
}

SEARCH_QUERIES = [
    'site:.ma "culture marocaine"',
    'site:.ma "patrimoine marocain"',
    'site:.ma "histoire du Maroc"',
    'site:.ma "traditions marocaines"',
    'site:.ma "darija marocaine"',
    'site:.ma "dialecte marocain"',
    'site:.ma "artisanat marocain"',
    'site:.ma "musique marocaine"',
    'site:.ma "cinema marocain"',
    'site:.ma "litterature marocaine"',
    'site:.ma "cuisine marocaine"',
    'site:.ma "architecture marocaine"',
    'site:.ma "musee maroc"',
    'site:.ma "archives maroc"',
    'site:.ma "Casablanca" culture',
    'site:.ma "Rabat" patrimoine',
    'site:.ma "Marrakech" traditions',
    'site:.ma "Fes" patrimoine',
    'site:.ma "Tanger" culture',
    'site:.ma "Souss" culture',
    'site:.ma "Sahara marocain" culture',
    'site:.ma المغرب الثقافة',
    'site:.ma التراث المغربي',
    'site:.ma التاريخ المغربي',
    'site:.ma الدارجة المغربية',
    'site:.ma العادات والتقاليد المغربية',
]

BANNED_DOMAINS = {
    "youtube.com",
    "youtu.be",
    "facebook.com",
    "fb.com",
    "instagram.com",
    "tiktok.com",
    "twitter.com",
    "x.com",
    "linkedin.com",
    "github.com",
    "gitlab.com",
    "kaggle.com",
    "huggingface.co",
    "zenodo.org",
    "wikipedia.org",
    "wikimedia.org",
    "google.com",
    "bing.com",
    "duckduckgo.com",
}

TRACKING_PARAMETERS = {
    "fbclid",
    "gclid",
    "mc_cid",
    "mc_eid",
    "ref",
    "source",
    "spm",
}

BANNED_PATH_FRAGMENTS = {
    "/login",
    "/logout",
    "/connexion",
    "/contact",
    "/user/",
    "/mon-compte",
    "/panier",
    "/qui-sommes-nous",
    "/organigramme",
    "/about/",
    "/publicite",
    "/abonnement",
    "/conditions-generales",
    "/politique-de-confidentialite",
    "/privacy",
    "/search",
    "/recherche",
    "/cdn-cgi/",
    "/meteo",
    "/horaire-priere",
    "/pharmacie-garde",
}

BANNED_EXTENSIONS = {
    ".jpg", ".jpeg", ".png", ".gif", ".webp", ".svg", ".ico",
    ".pdf", ".zip", ".rar", ".7z", ".xml", ".css", ".js",
    ".mp3", ".mp4", ".avi", ".mov", ".webm", ".woff", ".woff2",
}

MOROCCAN_SIGNALS = {
    "maroc",
    "marocain",
    "marocaine",
    "marocains",
    "morocco",
    "moroccan",
    "royaume du maroc",
    "darija",
    "maghreb",
    "casablanca",
    "rabat",
    "marrakech",
    "tanger",
    "agadir",
    "fes",
    "meknes",
    "oujda",
    "tetouan",
    "essaouira",
    "safi",
    "dakhla",
    "laayoune",
    "souss",
    "rif",
    "الأمازيغية",
    "الدارجة",
    "المغرب",
    "المغربي",
    "المغربية",
}

CATEGORY_SIGNALS = {
    "Darija": {"darija", "dialecte marocain", "arabe marocain", "الدارجة"},
    "Patrimoine": {"patrimoine", "heritage", "musee", "medina", "تراث", "متحف"},
    "Histoire": {"histoire", "historique", "archives", "dynastie", "تاريخ", "أرشيف"},
    "Traditions": {"tradition", "coutume", "festival", "moussem", "عادات", "تقاليد"},
    "Artisanat": {"artisanat", "artisan", "zellige", "tapis", "poterie", "صناعة تقليدية"},
    "Cuisine": {"cuisine", "recette", "tajine", "couscous", "gastronomie", "طبخ"},
    "Arts_et_medias": {"musique", "cinema", "film", "theatre", "litterature", "موسيقى", "سينما"},
    "Actualites_marocaines": {"actualite", "economie", "politique", "societe", "اخبار", "اقتصاد"},
}


@dataclass
class PageRecord:
    """Ligne finale exportee dans les trois formats de sortie."""

    url: str
    domaine: str
    titre: str
    description: str
    categorie: str
    score_marocain: int
    niveau_confiance: str
    source_decouverte: str
    date_collecte: str
    statut_http: int


def canonical_url(value: str, base_url: str = "") -> str:
    """Construit une URL HTTP propre et retire les parametres de tracking."""
    if not value:
        return ""
    absolute = urljoin(base_url, value.strip())
    parsed = urlparse(absolute)
    if parsed.scheme not in {"http", "https"} or not parsed.netloc:
        return ""
    query = [
        (key, val)
        for key, val in parse_qsl(parsed.query, keep_blank_values=False)
        if not key.lower().startswith("utm_") and key.lower() not in TRACKING_PARAMETERS
    ]
    path = re.sub(r"/{2,}", "/", parsed.path or "/")
    return urlunparse((parsed.scheme.lower(), parsed.netloc.lower(), path, "", urlencode(query), ""))


def domain_name(url: str) -> str:
    """Retourne le domaine sans le prefixe www."""
    return urlparse(url).netloc.lower().removeprefix("www.")


def is_banned(url: str) -> bool:
    """Exclut les reseaux sociaux, depots de code et catalogues de datasets."""
    domain = domain_name(url)
    path = urlparse(url).path.lower()
    banned_domain = any(domain == banned or domain.endswith(f".{banned}") for banned in BANNED_DOMAINS)
    banned_path = any(fragment in path for fragment in BANNED_PATH_FRAGMENTS)
    banned_extension = any(path.endswith(extension) for extension in BANNED_EXTENSIONS)
    return banned_domain or banned_path or banned_extension


def has_moroccan_site_identity(url: str) -> bool:
    """Evite de retenir un site international pour une seule page sur le Maroc."""
    domain = domain_name(url)
    return (
        domain.endswith(".ma")
        or domain in TRUSTED_SEED_DOMAINS
        or any(signal in domain for signal in ("maroc", "morocco", "bladi", "darija", "maghreb"))
    )


def visible_text(soup: BeautifulSoup) -> str:
    """Extrait un texte court utilisable pour valider le contexte marocain."""
    for tag in soup(["script", "style", "noscript", "svg"]):
        tag.decompose()
    return " ".join(soup.stripped_strings)[:50000]


def moroccan_score(url: str, title: str, description: str, text: str) -> int:
    """Calcule une confiance conservatrice sur l'identite marocaine du site."""
    domain = domain_name(url)
    haystack = f"{domain} {title} {description} {text}".lower()
    distinct_signals = sum(1 for signal in MOROCCAN_SIGNALS if signal in haystack)
    score = min(distinct_signals * 8, 64)
    if domain.endswith(".ma"):
        score += 30
    if any(word in domain for word in ("maroc", "morocco", "maghreb", "darija")):
        score += 20
    if "royaume du maroc" in haystack or "المملكة المغربية" in haystack:
        score += 15
    return min(score, 100)


def category_for(text: str) -> str:
    """Classe la page avec quelques categories marocaines faciles a relire."""
    normalized = text.lower()
    scores = {
        category: sum(1 for signal in signals if signal in normalized)
        for category, signals in CATEGORY_SIGNALS.items()
    }
    category, score = max(scores.items(), key=lambda item: item[1])
    return category if score else "Contexte_marocain"


def extract_result_urls(html: str) -> list[str]:
    """Extrait les URLs externes concretes d'une page de resultats Bing."""
    soup = BeautifulSoup(html, "html.parser")
    urls: list[str] = []
    for anchor in soup.select("li.b_algo h2 a[href], #b_results a[href]"):
        url = canonical_url(anchor.get("href", ""))
        if url and not is_banned(url):
            urls.append(url)
    return list(dict.fromkeys(urls))


def duckduckgo_target(href: str) -> str:
    """Recupere l'URL cible cachee dans un lien de redirection DuckDuckGo."""
    absolute = urljoin("https://duckduckgo.com", href)
    parsed = urlparse(absolute)
    target = (parse_qs(parsed.query).get("uddg") or [absolute])[0]
    return canonical_url(unquote(target))


async def discover_with_duckduckgo(queries: list[str], results_per_query: int) -> dict[str, str]:
    """Decouvre des sites marocains via le moteur HTML leger de DuckDuckGo."""
    found: dict[str, str] = {}
    headers = {"User-Agent": USER_AGENT, "Accept-Language": "fr-FR,ar-MA;q=0.8,ar;q=0.7"}
    async with httpx.AsyncClient(headers=headers, timeout=25, follow_redirects=True) as client:
        for index, query in enumerate(queries, start=1):
            LOGGER.info("Recherche web %s/%s: %s", index, len(queries), query)
            for offset in range(0, results_per_query, 10):
                try:
                    response = await client.get(
                        "https://html.duckduckgo.com/html/",
                        params={"q": query, "s": offset},
                    )
                    response.raise_for_status()
                except Exception as exc:
                    LOGGER.warning("Recherche DuckDuckGo interrompue pour %s: %s", query, exc)
                    break
                soup = BeautifulSoup(response.text, "html.parser")
                urls = [duckduckgo_target(anchor.get("href", "")) for anchor in soup.select(".result__a[href]")]
                urls = [url for url in urls if url and not is_banned(url)]
                before = len(found)
                for url in urls:
                    found.setdefault(url, query)
                if not urls or len(found) == before:
                    break
                await asyncio.sleep(0.4)
    return found


async def discover_with_playwright(queries: list[str], results_per_query: int) -> dict[str, str]:
    """Interroge Bing dans Chromium et associe chaque URL a sa requete source."""
    try:
        from playwright.async_api import async_playwright
    except ImportError as exc:
        raise RuntimeError("Installer Playwright avec: uv add playwright puis uv run playwright install chromium") from exc

    found: dict[str, str] = {}
    async with async_playwright() as playwright:
        browser = await playwright.chromium.launch(headless=True)
        context = await browser.new_context(
            user_agent=USER_AGENT,
            locale="fr-FR",
            extra_http_headers={"Accept-Language": "fr-FR,fr;q=0.9,ar-MA;q=0.8,ar;q=0.7"},
        )
        for index, query in enumerate(queries, start=1):
            LOGGER.info("Recherche %s/%s: %s", index, len(queries), query)
            # Une page neuve par requete isole les redirections tardives de Bing.
            page = await context.new_page()
            try:
                for offset in range(0, results_per_query, 10):
                    search_url = f"https://www.bing.com/search?q={quote_plus(query)}&first={offset + 1}"
                    await page.goto(search_url, wait_until="domcontentloaded", timeout=30000)
                    await page.wait_for_timeout(900)
                    anchors = page.locator("li.b_algo h2 a[href], #b_results a[href]")
                    await anchors.first.wait_for(state="attached", timeout=7000)
                    hrefs = await anchors.evaluate_all("elements => elements.map(element => element.href)")
                    urls = list(
                        dict.fromkeys(
                            url
                            for href in hrefs
                            if (url := canonical_url(href)) and not is_banned(url)
                        )
                    )
                    before = len(found)
                    for url in urls:
                        found.setdefault(url, query)
                    if len(found) == before:
                        break
            except Exception as exc:
                LOGGER.warning("Recherche interrompue pour %s: %s", query, exc)
            finally:
                await page.close()
        await context.close()
        await browser.close()
    return found


async def fetch_page(client: httpx.AsyncClient, url: str) -> tuple[str, int, str]:
    """Telecharge une page publique HTML et retourne URL finale, statut et contenu."""
    try:
        response = await client.get(url)
        content_type = response.headers.get("content-type", "").lower()
        if response.status_code >= 400 or "text/html" not in content_type:
            return url, response.status_code, ""
        return canonical_url(str(response.url)), response.status_code, response.text
    except Exception as exc:
        LOGGER.debug("Echec HTTP %s: %s", url, exc)
        return url, 0, ""


async def discover_sitemap_urls(
    client: httpx.AsyncClient,
    site_url: str,
    limit: int,
) -> list[str]:
    """Lit les sitemaps publics d'un domaine pour trouver rapidement ses pages."""
    parsed = urlparse(site_url)
    origin = f"{parsed.scheme}://{parsed.netloc}"
    sitemap_queue = [f"{origin}/sitemap.xml", f"{origin}/sitemap_index.xml"]
    discovered: list[str] = []
    visited_sitemaps: set[str] = set()

    try:
        robots = await client.get(f"{origin}/robots.txt")
        if robots.status_code < 400:
            for line in robots.text.splitlines():
                if line.lower().startswith("sitemap:"):
                    sitemap_queue.append(line.split(":", 1)[1].strip())
    except Exception:
        pass

    while sitemap_queue and len(discovered) < limit:
        sitemap_url = canonical_url(sitemap_queue.pop(0))
        if not sitemap_url or sitemap_url in visited_sitemaps:
            continue
        visited_sitemaps.add(sitemap_url)
        try:
            response = await client.get(sitemap_url)
            if response.status_code >= 400 or "<loc" not in response.text.lower():
                continue
        except Exception:
            continue
        soup = BeautifulSoup(response.text, "xml")
        for loc in soup.find_all("loc"):
            target = canonical_url(loc.get_text(strip=True))
            if not target or domain_name(target) != domain_name(site_url):
                continue
            if target.lower().endswith(".xml"):
                if len(visited_sitemaps) < 8:
                    sitemap_queue.append(target)
            elif not is_banned(target):
                discovered.append(target)
                if len(discovered) >= limit:
                    break
    return list(dict.fromkeys(discovered))


async def inspect_page(
    client: httpx.AsyncClient,
    url: str,
    source: str,
    minimum_score: int,
) -> tuple[PageRecord | None, list[str]]:
    """Valide le contexte marocain et extrait les liens internes de la page."""
    final_url, status, html = await fetch_page(client, url)
    if not html or is_banned(final_url):
        return None, []
    soup = BeautifulSoup(html, "html.parser")
    title = soup.title.get_text(" ", strip=True) if soup.title else ""
    meta = soup.select_one('meta[name="description"], meta[property="og:description"]')
    description = meta.get("content", "").strip() if meta else ""
    text = visible_text(soup)
    score = moroccan_score(final_url, title, description, text)
    if score < minimum_score or not has_moroccan_site_identity(final_url):
        return None, []

    internal: list[str] = []
    current_domain = domain_name(final_url)
    for anchor in soup.select("a[href]"):
        child = canonical_url(anchor.get("href", ""), final_url)
        if child and domain_name(child) == current_domain and not is_banned(child):
            internal.append(child)

    combined = f"{title} {description} {text[:10000]}"
    record = PageRecord(
        url=final_url,
        domaine=current_domain,
        titre=title[:500],
        description=description[:1000],
        categorie=category_for(combined),
        score_marocain=score,
        niveau_confiance="Tres_forte" if score >= 80 else "Forte",
        source_decouverte=source,
        date_collecte=datetime.now(timezone.utc).date().isoformat(),
        statut_http=status,
    )
    return record, list(dict.fromkeys(internal))


async def collect(args: argparse.Namespace) -> list[PageRecord]:
    """Orchestre la recherche, la validation et l'exploration interne."""
    discovered = {canonical_url(url): "seed_marocain" for url in SEED_SITES}
    if not args.skip_search:
        search_results = await discover_with_duckduckgo(SEARCH_QUERIES, args.results_per_query)
        # Le navigateur reste un secours lorsque le moteur HTML ne repond pas du tout.
        if not search_results:
            LOGGER.warning("DuckDuckGo n'a retourne aucun lien; activation du secours Playwright/Bing")
            search_results = await discover_with_playwright(SEARCH_QUERIES, args.results_per_query)
        discovered.update(search_results)
    LOGGER.info("%s URLs candidates trouvees", len(discovered))

    records: list[PageRecord] = []
    accepted_urls: set[str] = set()
    final_urls: set[str] = set()
    accepted_per_domain: dict[str, int] = {}
    queue: asyncio.Queue[tuple[str, str]] = asyncio.Queue()
    for url, source in discovered.items():
        if url and not is_banned(url):
            queue.put_nowait((url, source))

    headers = {"User-Agent": USER_AGENT, "Accept-Language": "fr-FR,ar-MA;q=0.8,ar;q=0.7"}
    limits = httpx.Limits(max_connections=args.concurrency, max_keepalive_connections=args.concurrency)
    async with httpx.AsyncClient(headers=headers, timeout=args.timeout, follow_redirects=True, limits=limits) as client:
        sitemap_results = await asyncio.gather(
            *(
                asyncio.wait_for(
                    discover_sitemap_urls(client, url, args.sitemap_links_per_site),
                    timeout=30,
                )
                for url in SEED_SITES
            ),
            return_exceptions=True,
        )
        sitemap_count = 0
        for seed_url, urls in zip(SEED_SITES, sitemap_results):
            if isinstance(urls, BaseException):
                LOGGER.debug("Sitemap ignore pour %s: %s", seed_url, urls)
                continue
            for url in urls:
                if url not in accepted_urls:
                    queue.put_nowait((url, f"sitemap:{domain_name(seed_url)}"))
                    sitemap_count += 1
        LOGGER.info("%s URLs ajoutees depuis les sitemaps", sitemap_count)

        next_checkpoint = 25
        while not queue.empty() and len(records) < args.max_links:
            batch: list[tuple[str, str]] = []
            while not queue.empty() and len(batch) < args.concurrency:
                url, source = queue.get_nowait()
                if accepted_per_domain.get(domain_name(url), 0) >= args.max_pages_per_site:
                    continue
                if url not in accepted_urls:
                    accepted_urls.add(url)
                    batch.append((url, source))
            if not batch:
                continue
            inspected = await asyncio.gather(
                *(inspect_page(client, url, source, args.minimum_score) for url, source in batch)
            )
            for (record, children), (_, source) in zip(inspected, batch):
                if record is None:
                    continue
                if record.url in final_urls:
                    continue
                domain_count = accepted_per_domain.get(record.domaine, 0)
                if domain_count >= args.max_pages_per_site:
                    continue
                records.append(record)
                final_urls.add(record.url)
                accepted_per_domain[record.domaine] = domain_count + 1
                LOGGER.info("[%s/%s] %s", len(records), args.max_links, record.url)
                if accepted_per_domain[record.domaine] < args.max_pages_per_site:
                    for child in children[: args.internal_links_per_page]:
                        if child not in accepted_urls:
                            queue.put_nowait((child, source))
                if len(records) >= args.max_links:
                    break
            if len(records) >= next_checkpoint:
                export_csv(records)
                export_jsonl(records)
                LOGGER.info("Sauvegarde intermediaire: %s pages", len(records))
                next_checkpoint += 25
    return records


def export_csv(records: list[PageRecord]) -> None:
    """Ecrit un CSV simple ouvrable avec Excel ou un autre projet Python."""
    fieldnames = list(PageRecord.__dataclass_fields__)
    with OUTPUT_CSV.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(asdict(record) for record in records)


def export_jsonl(records: list[PageRecord]) -> None:
    """Ecrit une ligne JSON par page pour faciliter les traitements automatiques."""
    with OUTPUT_JSONL.open("w", encoding="utf-8") as handle:
        for record in records:
            handle.write(json.dumps(asdict(record), ensure_ascii=False) + "\n")


def export_excel(records: list[PageRecord]) -> None:
    """Cree le livrable Excel avec une colonne cliquable pour chaque URL."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sites_marocains"
    columns = list(PageRecord.__dataclass_fields__) + ["ouvrir"]
    sheet.append(columns)
    for row_number, record in enumerate(records, start=2):
        values = list(asdict(record).values())
        values.append(f'=HYPERLINK(A{row_number},"Ouvrir")')
        sheet.append(values)

    fill = PatternFill("solid", fgColor="C1272D")
    for cell in sheet[1]:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = {"A": 65, "B": 28, "C": 50, "D": 70, "E": 24, "F": 18, "G": 18, "H": 40, "I": 16, "J": 14, "K": 12}
    for letter, width in widths.items():
        sheet.column_dimensions[letter].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    if records:
        table = Table(displayName="T_SitesMarocains", ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        sheet.add_table(table)
    workbook.save(OUTPUT_XLSX)


def build_parser() -> argparse.ArgumentParser:
    """Declare les quelques options utiles sans configuration externe."""
    parser = argparse.ArgumentParser(description="Collecteur simple de sites web a contenu marocain")
    parser.add_argument("--max-links", type=int, default=1000, help="Nombre maximal de pages finales")
    parser.add_argument("--results-per-query", type=int, default=50, help="Resultats recherches par requete")
    parser.add_argument("--max-pages-per-site", type=int, default=40, help="Pages internes conservees par site")
    parser.add_argument("--internal-links-per-page", type=int, default=30, help="Liens internes ajoutes par page")
    parser.add_argument("--minimum-score", type=int, default=55, help="Confiance marocaine minimale sur 100")
    parser.add_argument("--concurrency", type=int, default=10, help="Requetes HTTP simultanees")
    parser.add_argument("--timeout", type=float, default=20.0, help="Timeout HTTP en secondes")
    parser.add_argument(
        "--sitemap-links-per-site",
        type=int,
        default=300,
        help="URLs maximales ajoutees depuis les sitemaps de chaque site",
    )
    parser.add_argument(
        "--skip-search",
        action="store_true",
        help="Explore directement les sites marocains connus sans attendre les moteurs publics",
    )
    return parser


def main() -> None:
    """Lance la collecte puis ecrit les trois fichiers finaux a la racine."""
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
    args = build_parser().parse_args()
    records = asyncio.run(collect(args))
    export_csv(records)
    export_jsonl(records)
    export_excel(records)
    LOGGER.info("Termine: %s pages marocaines exportees", len(records))
    LOGGER.info("Excel: %s", OUTPUT_XLSX)
    LOGGER.info("CSV: %s", OUTPUT_CSV)
    LOGGER.info("JSONL: %s", OUTPUT_JSONL)


if __name__ == "__main__":
    main()
