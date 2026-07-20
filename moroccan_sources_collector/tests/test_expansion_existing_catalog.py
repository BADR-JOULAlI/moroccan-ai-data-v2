"""Tests automatis?s qui documentent et prot?gent le comportement attendu du projet."""

from pathlib import Path

from openpyxl import Workbook, load_workbook

from src.expansion.existing_catalog import load_existing_catalog
from src.storage.expansion_excel_exporter import EXPANDED_COLUMNS, export_expanded_excel


def test_load_existing_catalog_extracts_urls(tmp_path: Path):
    """V?rifie le comportement attendu d?une fonctionnalit? du projet."""
    path = tmp_path / "catalogue.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Catalogue_liens"
    ws.append(["URL_originale", "URL_finale", "Domaine", "Nom_source", "Priorite"])
    ws.append(["https://www.example.com/a?utm_source=x", "https://example.com/a", "example.com", "Example", "P1"])
    wb.save(path)
    catalog = load_existing_catalog(path)
    assert len(catalog.records) == 1
    assert "https://example.com/a" in catalog.normalized_urls
    assert "example.com" in catalog.domains


def test_expanded_excel_keeps_initial_and_new_rows(tmp_path: Path):
    """V?rifie le comportement attendu d?une fonctionnalit? du projet."""
    old = {column: "" for column in EXPANDED_COLUMNS}
    old.update({"ID": 1, "URL_finale": "https://example.com/old", "Est_nouveau": "Non", "Domaine": "example.com"})
    new = {column: "" for column in EXPANDED_COLUMNS}
    new.update({"ID": 2, "URL_finale": "https://new.example/new", "Est_nouveau": "Oui", "Domaine": "new.example", "Statut_lien": "Accessible"})
    path = tmp_path / "expanded.xlsx"
    export_expanded_excel([old, new], [new], [], path, 1)
    wb = load_workbook(path, data_only=False)
    assert "Nouveaux_liens" in wb.sheetnames
    ws = wb["Catalogue_liens"]
    assert ws.max_row == 3
    assert ws.cell(2, EXPANDED_COLUMNS.index("Ouvrir") + 1).value == '=HYPERLINK(D2,"Ouvrir")'
