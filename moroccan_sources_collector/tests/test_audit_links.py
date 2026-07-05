from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from src.audit.audit_exporter import export_audit_excel
from src.audit.github_validator import validate_github
from src.audit.huggingface_validator import validate_huggingface
from src.audit.replacement_extractor import extract_concrete_links
from src.audit.search_page_detector import concrete_type_for_url, is_search_page
from src.audit.soft_404_detector import detect_content_problem
from src.audit.url_auditor import AuditConfig, UrlAuditor
from src.audit.zenodo_validator import validate_zenodo
from src.utils.url_utils import url_fingerprint


def test_github_search_zero_result_is_rejected() -> None:
    html = "<html><title>Search</title><body>Your search did not match any repositories</body></html>"
    hint = validate_github("https://github.com/search?q=darija&type=repositories", "Search", html, html)

    assert hint["Type_URL"] == "PAGE_RECHERCHE"
    assert hint["Decision_hint"] == "REJETER"
    assert hint["Raison_hint"] == "ZERO_RESULTAT"


def test_url_auditor_decision_for_http_200_zero_result() -> None:
    auditor = UrlAuditor(Path("unused.json"), AuditConfig())
    check = {"status_http": 200, "status_lien": "Accessible"}
    problem = {"problem": "ZERO_RESULTAT", "message": "Your search did not match any repositories"}

    decision, reason = auditor._decision("PAGE_RECHERCHE", check, problem, {}, {"Score_pertinence": 0})

    assert decision == "REJETER"
    assert reason == "ZERO_RESULTAT"


def test_github_concrete_repository_type() -> None:
    assert concrete_type_for_url("https://github.com/owner/repository") == "PAGE_REPOSITORY"


def test_huggingface_search_and_dataset_detection() -> None:
    search_html = '<a href="/datasets/org/moroccan-darija">dataset</a>'
    hint = validate_huggingface("https://huggingface.co/datasets?search=darija", "", "", search_html)

    assert hint["Type_URL"] == "PAGE_RECHERCHE"
    assert hint["Decision_hint"] == "REMPLACER"
    assert "https://huggingface.co/datasets/org/moroccan-darija" in hint["URL_remplacement"]
    assert concrete_type_for_url("https://huggingface.co/datasets/org/moroccan-darija") == "PAGE_DATASET"


def test_zenodo_search_and_record_detection() -> None:
    search_html = '<a href="/records/123456">record</a>'
    hint = validate_zenodo("https://zenodo.org/search?q=darija", "", "", search_html)

    assert hint["Type_URL"] == "PAGE_RECHERCHE"
    assert hint["Decision_hint"] == "REMPLACER"
    assert "https://zenodo.org/records/123456" in hint["URL_remplacement"]
    assert concrete_type_for_url("https://zenodo.org/records/123456") == "PAGE_DATASET"


def test_soft_404_and_login_detection() -> None:
    assert detect_content_problem(text="Page not found")["problem"] == "SOFT_404"
    assert detect_content_problem(text="Sign in to continue")["problem"] == "PAGE_CONNEXION"


def test_search_page_detection_for_platforms() -> None:
    assert is_search_page("https://github.com/search?q=darija&type=repositories")
    assert is_search_page("https://huggingface.co/datasets?search=darija")
    assert is_search_page("https://youtube.com/results?search_query=darija")


def test_concrete_link_extraction_from_search_pages() -> None:
    github_html = '<a href="/open-darija/dataset">repo</a><a href="/topics/nlp">topic</a>'
    hf_html = '<a href="/datasets/org/darija-corpus">dataset</a>'

    assert extract_concrete_links("https://github.com/search?q=darija", github_html) == [
        "https://github.com/open-darija/dataset"
    ]
    assert extract_concrete_links("https://huggingface.co/datasets?search=darija", hf_html) == [
        "https://huggingface.co/datasets/org/darija-corpus"
    ]


def test_url_fingerprint_deduplicates_tracking_and_www() -> None:
    first = url_fingerprint("http://www.example.ma/page/?utm_source=test#x")
    second = url_fingerprint("https://example.ma/page")

    assert first == second


def test_export_excel_sheets_and_hyperlink_formula(tmp_path: Path) -> None:
    records = [
        {
            "ID": 1,
            "Nom_source": "Darija corpus",
            "URL_originale": "https://example.ma/darija",
            "URL_normalisee": "https://example.ma/darija",
            "URL_finale": "https://example.ma/darija",
            "URL_canonique": "https://example.ma/darija",
            "Domaine": "example.ma",
            "Plateforme": "example.ma",
            "Type_URL": "PAGE_CONTENU",
            "Statut_HTTP": 200,
            "Decision_finale": "CONSERVER",
            "Raison_decision": "RESSOURCE_CONCRETE_PERTINENTE",
            "Score_pertinence": 80,
        },
        {
            "ID": 2,
            "URL_originale": "https://github.com/search?q=none&type=repositories",
            "URL_normalisee": "https://github.com/search?q=none&type=repositories",
            "URL_finale": "https://github.com/search?q=none&type=repositories",
            "URL_canonique": "https://github.com/search?q=none&type=repositories",
            "Domaine": "github.com",
            "Plateforme": "GitHub",
            "Type_URL": "PAGE_RECHERCHE",
            "Statut_HTTP": 200,
            "Message_erreur_detecte": "Your search did not match any repositories",
            "Decision_finale": "REJETER",
            "Raison_decision": "ZERO_RESULTAT",
            "Score_pertinence": 0,
        },
    ]
    output = tmp_path / "audit.xlsx"

    export_audit_excel(records, [], output, total_loaded=2)

    wb = load_workbook(output, data_only=False)
    assert {
        "Catalogue_valide",
        "Liens_remplaces",
        "Pages_recherche",
        "Liens_rejetes",
        "Liens_a_recontroler",
        "Doublons",
        "Rapport_audit",
        "Tableau_de_bord",
    }.issubset(set(wb.sheetnames))
    assert wb["Catalogue_valide"]["AF2"].value == '=HYPERLINK(E2,"Ouvrir")'
    valid_urls = [row[0] for row in wb["Catalogue_valide"].iter_rows(min_row=2, min_col=3, max_col=3, values_only=True)]
    assert "https://github.com/search?q=none&type=repositories" not in valid_urls
