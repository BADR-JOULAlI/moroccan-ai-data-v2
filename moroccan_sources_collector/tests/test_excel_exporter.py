"""Tests automatis?s qui documentent et prot?gent le comportement attendu du projet."""

from pathlib import Path

from openpyxl import load_workbook

from src.storage.excel_exporter import CATALOG_COLUMNS, export_excel


def test_excel_exporter_creates_sheets_and_hyperlink_formula(tmp_path: Path):
    """V?rifie le comportement attendu d?une fonctionnalit? du projet."""
    record = {column: "" for column in CATALOG_COLUMNS}
    record.update({
        "ID": 1,
        "Nom_source": "Test",
        "URL_originale": "https://example.com",
        "URL_finale": "https://example.com",
        "Domaine": "example.com",
        "Categorie_principale": "Darija",
        "Type_contenu": "HTML",
        "Langue_principale": "darija_latine",
        "Contenu_parallele": "Oui",
        "Statut_lien": "Accessible",
        "Licence": "Inconnue",
        "Priorite": "P0",
    })
    path = tmp_path / "out.xlsx"
    export_excel([record], path)
    wb = load_workbook(path, data_only=False)
    assert set(["Catalogue_liens", "Sources_prioritaires", "Liens_a_recontroler", "Sources_hors_perimetre", "Sources_francais_darija", "Tableau_de_bord", "Referentiel"]).issubset(wb.sheetnames)
    ws = wb["Catalogue_liens"]
    assert ws.cell(2, CATALOG_COLUMNS.index("Ouvrir") + 1).value == '=HYPERLINK(D2,"Ouvrir")'
