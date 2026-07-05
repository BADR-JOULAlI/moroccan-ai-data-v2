from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
REPORTS = ROOT / "reports"
OUTPUT = REPORTS / "rapport_modifications_aujourdhui.docx"

BLUE = "1F4E78"
LIGHT_BLUE = "E8EEF5"
GREEN = "D9EAD3"
GRAY = "F2F4F7"
DARK = "0B2545"


def set_cell_fill(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(table, top=80, start=120, bottom=80, end=120):
    tbl_pr = table._tbl.tblPr
    margins = tbl_pr.first_child_found_in("w:tblCellMar")
    if margins is None:
        margins = OxmlElement("w:tblCellMar")
        tbl_pr.append(margins)
    for key, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = margins.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(11)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].paragraph_format.line_spacing = 1.10
    for name, size, color in [
        ("Heading 1", 16, BLUE),
        ("Heading 2", 13, BLUE),
        ("Heading 3", 12, DARK),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)


def add_title(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Rapport des modifications realisees aujourd'hui")
    run.bold = True
    run.font.size = Pt(23)
    run.font.color.rgb = RGBColor.from_string(DARK)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Projet de catalogue de sources pour une IA marocaine")
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor.from_string("555555")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"Date : {datetime.now().strftime('%d/%m/%Y')}")
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor.from_string("666666")


def add_callout(doc: Document, title: str, text: str, fill: str = LIGHT_BLUE) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_cell_margins(table)
    cell = table.cell(0, 0)
    set_cell_fill(cell, fill)
    p = cell.paragraphs[0]
    r = p.add_run(title)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(BLUE)
    p = cell.add_paragraph(text)
    p.paragraph_format.space_after = Pt(0)
    doc.add_paragraph()


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def add_summary_table(doc: Document) -> None:
    data = [
        ("Fichier Excel final", "catalogue_sources_marocaines_ai_FINAL.xlsx"),
        ("Lignes de donnees", "727"),
        ("Nouveaux liens ajoutes", "500"),
        ("Liens YouTube ajoutes", "40"),
        ("Liens Facebook ajoutes", "40"),
        ("Tests automatises", "26 passed"),
        ("Evaluation indicative Darija", "81.67% sur 60 exemples"),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    set_cell_margins(table)
    headers = table.rows[0].cells
    headers[0].text = "Indicateur"
    headers[1].text = "Resultat"
    for cell in headers:
        set_cell_fill(cell, LIGHT_BLUE)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
    for label, value in data:
        row = table.add_row().cells
        row[0].text = label
        row[1].text = value


def build() -> Path:
    REPORTS.mkdir(parents=True, exist_ok=True)
    doc = Document()
    style_document(doc)
    add_title(doc)

    add_callout(
        doc,
        "Synthese",
        "Aujourd'hui, j'ai enrichi le catalogue de sources, produit un fichier Excel final clair, "
        "et ameliore le detecteur de Darija afin de rendre la detection plus precise, explicable et maintenable.",
        GREEN,
    )

    doc.add_heading("1. Enrichissement du catalogue", level=1)
    doc.add_paragraph(
        "J'ai ajoute de nouvelles strategies de decouverte de sources, notamment l'utilisation de liens publics YouTube et Facebook "
        "lorsqu'ils sont en relation avec les mots-cles du projet : Darija, culture marocaine, arabe marocain, corpus, datasets, videos et contenus publics pertinents."
    )
    add_bullets(doc, [
        "Generation d'un fichier Excel final non ambigu : catalogue_sources_marocaines_ai_FINAL.xlsx.",
        "Conservation des lignes existantes et ajout des nouveaux liens dans des feuilles separees.",
        "Verification du classeur avec openpyxl : feuilles, lignes, colonnes et formules HYPERLINK.",
        "Production de 727 lignes de donnees dans la feuille Catalogue_liens.",
    ])

    doc.add_heading("2. Amelioration du detecteur de Darija", level=1)
    doc.add_paragraph(
        "J'ai remplace l'ancien detecteur base sur une petite liste codee en dur par un detecteur plus robuste, base sur un lexique YAML pondere."
    )
    add_bullets(doc, [
        "Creation de config/darija_lexicon.yaml pour externaliser les mots, expressions, variantes et indicateurs marocains.",
        "Ajout de poids strong, medium et weak pour eviter qu'un mot faible seul classe un texte comme Darija.",
        "Detection des expressions composees de deux a cinq mots.",
        "Ajout de sous-scores explicables : lexical, expressions, contexte marocain et penalite d'ambiguite.",
    ])

    doc.add_heading("3. Correction de l'Arabizi", level=1)
    doc.add_paragraph(
        "J'ai corrige la detection Arabizi afin de ne plus confondre les annees ou les tokens techniques avec de vrais mots Arabizi."
    )
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    set_cell_margins(table)
    for i, h in enumerate(["Exemple", "Ancien risque", "Nouveau resultat"]):
        table.rows[0].cells[i].text = h
        set_cell_fill(table.rows[0].cells[i], LIGHT_BLUE)
    rows = [
        ("2024", "Faux Arabizi possible", "Non detecte comme Arabizi"),
        ("Python3 / HTTP2 / UTF8", "Faux positif technique", "Ignore correctement"),
        ("sha256 / ipv4 / html5", "Identifiant technique confondu", "Ignore correctement"),
        ("3lach / 7na / b7al", "Vrai Arabizi", "Detecte correctement"),
    ]
    for row_data in rows:
        row = table.add_row().cells
        for i, value in enumerate(row_data):
            row[i].text = value

    doc.add_heading("4. Tests et evaluation", level=1)
    add_bullets(doc, [
        "Ajout et mise a jour des tests unitaires du detecteur.",
        "Validation des cas Darija arabe, Darija latine, Arabizi, arabe standard, mots faibles, repetition, texte vide et texte long.",
        "Creation d'un jeu de validation local de 60 exemples.",
        "Creation du script scripts/evaluate_darija_detector.py.",
    ])
    add_summary_table(doc)

    doc.add_heading("5. Documentation et preparation future", level=1)
    doc.add_paragraph(
        "J'ai mis a jour le README avec une section expliquant le fonctionnement du detecteur de Darija, les seuils, la normalisation, "
        "l'ajout de mots, de variantes et d'expressions. J'ai aussi prepare une architecture experimentale pour une future detection supervisee."
    )

    doc.add_heading("6. Conclusion", level=1)
    doc.add_paragraph(
        "Les modifications realisees aujourd'hui renforcent le projet sur deux aspects essentiels : le volume exploitable du catalogue et la qualite de l'analyse linguistique. "
        "Le projet dispose maintenant d'un fichier Excel final clair, d'un detecteur Darija plus maintenable, d'une meilleure detection Arabizi et d'une base de tests plus solide."
    )

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Rapport interne - Projet IA marocaine")

    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
