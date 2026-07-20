"""Module d?audit technique des URLs et des fichiers Excel g?n?r?s."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

URL_ALIASES = ["URL", "URL_originale", "URL_finale", "URL_canonique", "Link", "Lien"]
PREFERRED_SHEETS = ["Catalogue_liens", "Nouveaux_liens", "Sources", "Sheet1"]


def resolve_excel_path(input_value: str, root: Path) -> Path:
    """R?sout une valeur d?entr?e vers une ressource concr?te utilisable par le pipeline."""
    path = Path(input_value)
    candidates = [
        path,
        root / input_value,
        root / "data" / "output" / input_value,
        root / "data" / "output" / "catalogue_sources_marocaines_ai_FINAL.xlsx",
    ]
    for candidate in candidates:
        if candidate.exists():
            return candidate
    raise FileNotFoundError(f"Excel file not found: {input_value}")


def _normalize_header(value: Any) -> str:
    """D?crit le fonctionnement de `_normalize_header` et retourne le r?sultat attendu par le module."""
    return str(value or "").strip().lower().replace(" ", "_")


def _select_sheet(wb):
    """D?crit le fonctionnement de `_select_sheet` et retourne le r?sultat attendu par le module."""
    for name in PREFERRED_SHEETS:
        if name in wb.sheetnames:
            return wb[name]
    return wb[wb.sheetnames[0]]


def load_url_rows(path: Path) -> tuple[list[dict], str, str]:
    """Charge des donn?es ou une configuration depuis le disque et retourne une structure exploitable."""
    wb = load_workbook(path, data_only=False, read_only=True)
    ws = _select_sheet(wb)
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], ws.title, ""
    headers = [str(h or "").strip() for h in rows[0]]
    normalized = {_normalize_header(h): h for h in headers}
    url_col = ""
    for alias in URL_ALIASES:
        key = _normalize_header(alias)
        if key in normalized:
            url_col = normalized[key]
            break
    if not url_col:
        raise ValueError(f"No URL column found in sheet {ws.title}")
    records: list[dict] = []
    for idx, row in enumerate(rows[1:], start=2):
        record = {headers[i]: row[i] if i < len(row) else "" for i in range(len(headers))}
        url = str(record.get(url_col) or "").strip()
        if not url:
            continue
        record["_row_number"] = idx
        record["_input_url_column"] = url_col
        record["_input_url"] = url
        records.append(record)
    return records, ws.title, url_col
