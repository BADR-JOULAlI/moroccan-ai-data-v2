from __future__ import annotations

from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

AUDIT_COLUMNS = [
    "ID", "Nom_source", "URL_originale", "URL_normalisee", "URL_finale", "URL_canonique", "Domaine", "Plateforme",
    "Type_URL", "Statut_HTTP", "Type_MIME", "Titre_page", "Meta_description", "Message_erreur_detecte",
    "Nombre_resultats", "Est_page_recherche", "Est_soft_404", "Est_page_connexion", "JavaScript_requis",
    "Robots_txt", "Categorie_principale", "Langue_principale", "Score_darija", "Score_culture_marocaine",
    "Score_francais_darija", "Score_pertinence", "Decision_finale", "Raison_decision", "URL_remplacement",
    "Date_verification", "Notes", "Ouvrir",
]


def _format(ws, name: str) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows(min_row=2):
        values = {ws.cell(1, idx + 1).value: row[idx].value for idx in range(len(row))}
        color = None
        if values.get("Decision_finale") == "CONSERVER":
            color = "D9EAD3"
        elif values.get("Decision_finale") == "REJETER":
            color = "F4CCCC"
        elif values.get("Decision_finale") == "A_RECONTROLER":
            color = "FFF2CC"
        if color:
            fill_row = PatternFill("solid", fgColor=color)
            for cell in row:
                cell.fill = fill_row
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for col in ws.columns:
        width = min(max(max(len(str(c.value or "")) for c in col[:200]) + 2, 10), 60)
        ws.column_dimensions[col[0].column_letter].width = width
    if ws.max_row >= 2:
        table = Table(displayName=f"T_{name[:24]}", ref=ws.dimensions)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)


def _write(wb: Workbook, title: str, records: list[dict], columns: list[str] = AUDIT_COLUMNS) -> None:
    ws = wb.create_sheet(title)
    ws.append(columns)
    for record in records:
        row_num = ws.max_row + 1
        ws.append([f'=HYPERLINK(E{row_num},"Ouvrir")' if col == "Ouvrir" else record.get(col, "") for col in columns])
    _format(ws, title)


def _dashboard(wb: Workbook, records: list[dict]) -> None:
    ws = wb.create_sheet("Tableau_de_bord")
    ws.append(["Indicateur", "Valeur"])
    metrics = [
        ("Total audite", len(records)),
        ("Sources conservees", sum(r["Decision_finale"] == "CONSERVER" for r in records)),
        ("Pages de recherche", sum(r["Type_URL"] == "PAGE_RECHERCHE" for r in records)),
        ("Soft 404", sum(r["Type_URL"] == "SOFT_404" for r in records)),
        ("Pages connexion", sum(r["Type_URL"] == "PAGE_CONNEXION" for r in records)),
        ("A recontroler", sum(r["Decision_finale"] == "A_RECONTROLER" for r in records)),
        ("Rejetees", sum(r["Decision_finale"] == "REJETER" for r in records)),
    ]
    for row in metrics:
        ws.append(list(row))
    start = ws.max_row + 2
    for label, field in [("Decisions finales", "Decision_finale"), ("Types URL", "Type_URL"), ("Plateformes", "Plateforme"), ("Erreurs", "Raison_decision"), ("Categories", "Categorie_principale")]:
        ws.cell(start, 1, label)
        ws.cell(start, 2, "Nombre")
        counts = Counter(r.get(field) or "Inconnu" for r in records)
        row = start + 1
        for key, value in counts.most_common(20):
            ws.cell(row, 1, key)
            ws.cell(row, 2, value)
            row += 1
        chart = PieChart() if field == "Decision_finale" else BarChart()
        chart.title = label
        chart.add_data(Reference(ws, min_col=2, min_row=start, max_row=row - 1), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=start + 1, max_row=row - 1))
        ws.add_chart(chart, f"D{start}")
        start = row + 15
    _format(ws, "Dashboard")


def export_audit_excel(records: list[dict], duplicates: list[dict], output_path: Path, total_loaded: int) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    valid = [r for r in records if r["Decision_finale"] == "CONSERVER"]
    replaced = [r for r in records if r["Decision_finale"] == "REMPLACER"]
    searches = [r for r in records if r["Type_URL"] == "PAGE_RECHERCHE"]
    rejected = [r for r in records if r["Decision_finale"] == "REJETER"]
    recheck = [r for r in records if r["Decision_finale"] == "A_RECONTROLER"]
    _write(wb, "Catalogue_valide", valid)
    _write(wb, "Liens_remplaces", replaced)
    _write(wb, "Pages_recherche", searches)
    _write(wb, "Liens_rejetes", rejected)
    _write(wb, "Liens_a_recontroler", recheck)
    _write(wb, "Doublons", duplicates, ["URL_candidate", "URL_existante", "type_doublon", "raison", "score_similarite"])
    report_rows = [
        {"Indicateur": "nombre total de liens lus", "Valeur": total_loaded},
        {"Indicateur": "nombre de liens accessibles", "Valeur": sum(str(r.get("Statut_HTTP")).startswith("2") for r in records)},
        {"Indicateur": "nombre de vraies ressources", "Valeur": len(valid)},
        {"Indicateur": "nombre de pages de recherche", "Valeur": len(searches)},
        {"Indicateur": "nombre de recherches avec zero resultat", "Valeur": sum(r["Raison_decision"] == "ZERO_RESULTAT" for r in records)},
        {"Indicateur": "nombre de soft 404", "Valeur": sum(r["Type_URL"] == "SOFT_404" for r in records)},
        {"Indicateur": "nombre de liens morts", "Valeur": sum(r["Statut_HTTP"] == 404 for r in records)},
        {"Indicateur": "nombre de pages de connexion", "Valeur": sum(r["Type_URL"] == "PAGE_CONNEXION" for r in records)},
        {"Indicateur": "nombre de liens remplaces", "Valeur": len(replaced)},
        {"Indicateur": "nombre de nouvelles ressources extraites", "Valeur": sum(1 for r in records if r.get("URL_recherche_origine"))},
        {"Indicateur": "nombre de doublons", "Valeur": len(duplicates)},
        {"Indicateur": "nombre de liens rejetes", "Valeur": len(rejected)},
        {"Indicateur": "nombre de liens a recontroler", "Valeur": len(recheck)},
        {"Indicateur": "nombre final de sources valides", "Valeur": len(valid)},
    ]
    _write(wb, "Rapport_audit", report_rows, ["Indicateur", "Valeur"])
    _dashboard(wb, records)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
