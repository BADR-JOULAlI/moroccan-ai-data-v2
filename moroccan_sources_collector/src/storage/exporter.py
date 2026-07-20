"""Exporteur final du MVP avec un classeur Excel segmente par plateforme."""

from __future__ import annotations

import re
from pathlib import Path
from urllib.parse import urlparse

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from src.storage.excel_exporter import CATALOG_COLUMNS


class DataExporter:
    """Convertit les enregistrements Python valides en fichiers finaux du MVP.

    La classe recoit une liste globale de dictionnaires. Chaque dictionnaire
    represente une source valide du catalogue final. L'export Excel ajoute une
    feuille globale, puis une feuille par plateforme detectee: YouTube, Facebook,
    Kaggle, GitHub ou autre plateforme future.
    """

    def __init__(self, records: list[dict]):
        """Memorise les enregistrements a exporter sans modifier la liste source."""
        self.records = [dict(record) for record in records]

    def export_excel_by_platform(self, path: Path) -> None:
        """Ecrit un classeur Excel multi-feuilles segmente par plateforme.

        Inputs:
            path: chemin du fichier Excel final destine au Ministere.

        Output:
            un fichier .xlsx contenant l'onglet ``Catalogue_liens`` et un onglet
            dedie pour chaque plateforme. Dans chaque feuille, la colonne
            ``Ouvrir`` contient une formule Excel ``HYPERLINK`` qui pointe vers
            ``URL_finale`` ou ``URL_originale``.
        """
        path.parent.mkdir(parents=True, exist_ok=True)
        prepared = self._records_with_platform_and_links(self.records)
        platform_groups = self._group_by_platform(prepared)

        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            self._write_dataframe(writer, "Catalogue_liens", prepared)
            for platform, records in platform_groups.items():
                self._write_dataframe(writer, self._sheet_name(platform), records)
            self._write_reference_sheet(writer, platform_groups)

        self._format_workbook(path)

    def _records_with_platform_and_links(self, records: list[dict]) -> list[dict]:
        """Ajoute les champs techniques necessaires a l'export Excel segmente."""
        prepared: list[dict] = []
        for index, record in enumerate(records, start=1):
            item = {column: record.get(column, "") for column in CATALOG_COLUMNS}
            for key, value in record.items():
                item.setdefault(key, value)
            item["ID"] = item.get("ID") or index
            item["platform"] = self._detect_platform(record)
            item["Ouvrir"] = self._hyperlink_formula(index + 1, item)
            prepared.append(item)
        return prepared

    def _group_by_platform(self, records: list[dict]) -> dict[str, list[dict]]:
        """Regroupe les lignes finales par plateforme pour creer les onglets."""
        grouped: dict[str, list[dict]] = {}
        for record in records:
            platform = record.get("platform") or "Autres"
            grouped.setdefault(str(platform), []).append(record)
        return dict(sorted(grouped.items(), key=lambda item: item[0].lower()))

    def _write_dataframe(self, writer: pd.ExcelWriter, sheet_name: str, records: list[dict]) -> None:
        """Transforme une liste de dictionnaires en feuille Excel pandas."""
        columns = self._ordered_columns(records)
        dataframe = pd.DataFrame(records, columns=columns)
        dataframe.to_excel(writer, sheet_name=sheet_name, index=False)

    def _write_reference_sheet(self, writer: pd.ExcelWriter, platform_groups: dict[str, list[dict]]) -> None:
        """Ajoute un onglet de synthese court pour expliquer la segmentation."""
        rows = [
            {
                "Plateforme": platform,
                "Nombre_liens": len(records),
                "Description": "Onglet genere automatiquement depuis la cle platform.",
            }
            for platform, records in platform_groups.items()
        ]
        pd.DataFrame(rows, columns=["Plateforme", "Nombre_liens", "Description"]).to_excel(
            writer,
            sheet_name="Synthese_plateformes",
            index=False,
        )

    def _ordered_columns(self, records: list[dict]) -> list[str]:
        """Place les colonnes ministerielles en premier puis conserve les extras."""
        extras = sorted({key for record in records for key in record if key not in CATALOG_COLUMNS})
        return CATALOG_COLUMNS + extras

    def _hyperlink_formula(self, excel_row: int, record: dict) -> str:
        """Construit la formule Excel qui ouvre directement la source."""
        url = record.get("URL_finale") or record.get("URL_originale") or record.get("url", "")
        if not url:
            return ""
        escaped = str(url).replace('"', '""')
        return f'=HYPERLINK("{escaped}","Ouvrir")'

    def _detect_platform(self, record: dict) -> str:
        """Identifie la plateforme a partir du champ explicite ou du domaine."""
        explicit = str(record.get("platform") or record.get("Platform") or "").strip()
        if explicit:
            return self._normalize_platform_name(explicit)
        category = str(record.get("Sous_categories") or record.get("Type_source") or "").lower()
        url = str(record.get("URL_finale") or record.get("URL_originale") or record.get("url") or "")
        domain = urlparse(url).netloc.lower()
        signal = f"{category} {domain} {url.lower()}"
        if "youtube" in signal or "youtu.be" in signal:
            return "YouTube"
        if "facebook" in signal or "fb.watch" in signal:
            return "Facebook"
        if "kaggle" in signal:
            return "Kaggle"
        if "github" in signal:
            return "GitHub"
        return "Autres"

    def _normalize_platform_name(self, value: str) -> str:
        """Uniformise les variantes de noms de plateformes avant groupement."""
        lowered = value.lower()
        if "youtube" in lowered:
            return "YouTube"
        if "facebook" in lowered:
            return "Facebook"
        if "kaggle" in lowered:
            return "Kaggle"
        if "github" in lowered:
            return "GitHub"
        return value[:1].upper() + value[1:]

    def _sheet_name(self, platform: str) -> str:
        """Nettoie le nom d'onglet pour respecter les contraintes Excel."""
        clean = re.sub(r"[\[\]\:\*\?\/\\]", "_", platform).strip() or "Autres"
        return clean[:31]

    def _format_workbook(self, path: Path) -> None:
        """Applique un format lisible sur toutes les feuilles generees."""
        from openpyxl import load_workbook

        workbook = load_workbook(path)
        header_fill = PatternFill("solid", fgColor="1F4E78")
        for sheet in workbook.worksheets:
            if sheet.max_row < 1:
                continue
            for cell in sheet[1]:
                cell.font = Font(color="FFFFFF", bold=True)
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for column in sheet.columns:
                letter = column[0].column_letter
                width = max(len(str(cell.value or "")) for cell in column[:100])
                sheet.column_dimensions[letter].width = min(max(width + 2, 10), 55)
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
            if sheet.max_row >= 2 and sheet.max_column >= 1:
                ref = f"A1:{sheet.cell(sheet.max_row, sheet.max_column).coordinate}"
                table_name = f"T_{re.sub(r'[^A-Za-z0-9_]', '_', sheet.title)[:25]}"
                table = Table(displayName=table_name, ref=ref)
                table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
                sheet.add_table(table)
        workbook.save(path)
