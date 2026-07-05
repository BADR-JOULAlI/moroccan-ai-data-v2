from __future__ import annotations

from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

CATALOG_COLUMNS = [
    "ID", "Nom_source", "URL_originale", "URL_finale", "Domaine", "Categorie_principale",
    "Sous_categories", "Description_source", "Type_source", "Type_contenu", "Langue_principale",
    "Pourcentage_script_arabe", "Pourcentage_script_latin", "Score_darija", "Mots_darija_detectes",
    "Arabizi_detecte", "Decision_linguistique", "Contenu_parallele", "Type_alignement",
    "Nombre_paires_detectees", "Exemples_paires", "Score_alignement", "Region_marocaine", "Usage_IA",
    "Methode_acces", "Statut_HTTP", "Statut_lien", "Redirection", "Robots_txt", "Licence",
    "URL_licence", "Autorisation_entrainement_IA", "Priorite", "Score_pertinence", "Score_fiabilite",
    "Date_verification", "Source_decouverte", "Mots_cles_associes", "Notes", "Ouvrir",
]


def _write_sheet(wb: Workbook, title: str, records: list[dict], columns: list[str]) -> None:
    ws = wb.create_sheet(title)
    ws.append(columns)
    for record in records:
        row = []
        for column in columns:
            if column == "Ouvrir":
                row.append(f'=HYPERLINK(D{ws.max_row + 1},"Ouvrir")')
            else:
                row.append(record.get(column, ""))
        ws.append(row)
    _format_table(ws, title.replace(" ", "_")[:25])


def _format_table(ws, table_name: str) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows(min_row=2):
        priority = row[32].value if len(row) > 32 else ""
        status = row[26].value if len(row) > 26 else ""
        fill = None
        if priority == "P0":
            fill = PatternFill("solid", fgColor="D9EAD3")
        elif priority == "P1":
            fill = PatternFill("solid", fgColor="FFF2CC")
        elif status in {"Timeout", "Bloque", "Erreur_SSL", "JavaScript_requis"}:
            fill = PatternFill("solid", fgColor="F4CCCC")
        if fill:
            for cell in row:
                cell.fill = fill
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for col in ws.columns:
        letter = col[0].column_letter
        max_len = max(len(str(cell.value or "")) for cell in col[:100])
        ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 55)
    if ws.max_row >= 2 and ws.max_column >= 1:
        ref = f"A1:{ws.cell(ws.max_row, ws.max_column).coordinate}"
        table = Table(displayName=f"T_{table_name}", ref=ref)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)


def _dashboard(wb: Workbook, records: list[dict]) -> None:
    ws = wb.create_sheet("Tableau_de_bord")
    domains = {record.get("Domaine") for record in records if record.get("Domaine")}
    metrics = [
        ("Nombre total de liens", len(records)),
        ("Domaines uniques", len(domains)),
        ("Liens accessibles", sum(r.get("Statut_lien") in {"Accessible", "Accessible_avec_redirection"} for r in records)),
        ("Liens a recontroler", sum(r.get("Statut_lien") not in {"Accessible", "Accessible_avec_redirection"} or r.get("Licence") == "Inconnue" for r in records)),
        ("Sources P0", sum(r.get("Priorite") == "P0" for r in records)),
        ("Sources P1", sum(r.get("Priorite") == "P1" for r in records)),
        ("Sources P2", sum(r.get("Priorite") == "P2" for r in records)),
        ("Sources Darija arabe", sum(r.get("Langue_principale") == "darija_arabe" for r in records)),
        ("Sources Darija latine", sum(r.get("Langue_principale") == "darija_latine" for r in records)),
        ("Sources Arabizi", sum(r.get("Arabizi_detecte") == "Oui" for r in records)),
        ("Sources arabe standard", sum(r.get("Langue_principale") == "arabe_standard" for r in records)),
        ("Sources francais-Darija", sum(r.get("Contenu_parallele") == "Oui" for r in records)),
        ("Sources exclues", sum(r.get("Decision_linguistique") == "Exclure" for r in records)),
    ]
    ws.append(["Indicateur", "Valeur"])
    for item in metrics:
        ws.append(list(item))
    start = ws.max_row + 2
    for title, field in [("Categorie", "Categorie_principale"), ("Langue", "Langue_principale"), ("Type contenu", "Type_contenu"), ("Statut", "Statut_lien")]:
        ws.cell(start, 1, title)
        ws.cell(start, 2, "Nombre")
        counts = Counter(r.get(field) or "Inconnu" for r in records)
        row = start + 1
        for key, value in counts.items():
            ws.cell(row, 1, key)
            ws.cell(row, 2, value)
            row += 1
        chart = PieChart() if title == "Langue" else BarChart()
        chart.title = f"Repartition par {title.lower()}"
        data = Reference(ws, min_col=2, min_row=start, max_row=row - 1)
        cats = Reference(ws, min_col=1, min_row=start + 1, max_row=row - 1)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        ws.add_chart(chart, f"D{start}")
        start = row + 16
    _format_table(ws, "Dashboard")


def _referential(wb: Workbook) -> None:
    ws = wb.create_sheet("Referentiel")
    rows = [
        ("Langues", "arabe_standard, darija_arabe, darija_latine, darija_arabizi, arabe_darija_mixte, francais_darija, autre_langue, indeterminee"),
        ("Decisions", "Conserver, A_recontroler, Exclure"),
        ("Priorites", "P0: critique; P1: pertinent; P2: complementaire"),
        ("Score_pertinence", "+Darija, +parallel, +dataset, +source officielle, +licence; penalites hors sujet et instabilite"),
        ("Robots/licence", "Une page publique n'est jamais consideree automatiquement libre pour l'entrainement IA."),
    ]
    ws.append(["Element", "Definition"])
    for row in rows:
        ws.append(row)
    _format_table(ws, "Referentiel")


def export_excel(records: list[dict], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    _write_sheet(wb, "Catalogue_liens", records, CATALOG_COLUMNS)
    _write_sheet(wb, "Sources_prioritaires", [r for r in records if r.get("Priorite") == "P0"], CATALOG_COLUMNS)
    recheck = [r for r in records if r.get("Statut_lien") not in {"Accessible", "Accessible_avec_redirection"} or r.get("Licence") == "Inconnue"]
    _write_sheet(wb, "Liens_a_recontroler", recheck, CATALOG_COLUMNS)
    out_scope = [r for r in records if r.get("Decision_linguistique") == "Exclure" or r.get("Langue_principale") == "autre_langue"]
    _write_sheet(wb, "Sources_hors_perimetre", out_scope, CATALOG_COLUMNS)
    fr_darija = [r for r in records if r.get("Contenu_parallele") == "Oui" or "francais_darija" in str(r.get("Categorie_principale"))]
    _write_sheet(wb, "Sources_francais_darija", fr_darija, CATALOG_COLUMNS)
    _dashboard(wb, records)
    _referential(wb)
    wb.save(path)
