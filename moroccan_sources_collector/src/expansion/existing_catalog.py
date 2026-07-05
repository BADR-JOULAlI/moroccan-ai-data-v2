from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from src.utils.url_utils import get_domain, normalize_url, url_fingerprint


URL_COLUMNS = {"url_originale", "url_finale", "url"}


@dataclass
class ExistingCatalog:
    records: list[dict]
    original_urls: set[str]
    final_urls: set[str]
    normalized_urls: set[str]
    domains: set[str]
    fingerprints: set[str]


def resolve_input_path(input_value: str, root: Path) -> Path:
    path = Path(input_value)
    if path.exists():
        return path
    candidates = [
        root / input_value,
        root / "data" / "output" / input_value,
    ]
    for candidate in candidates:
        if candidate.exists():
            return candidate
    raise FileNotFoundError(f"Input Excel not found: {input_value}")


def _header_map(headers: list[str]) -> dict[str, str]:
    normalized = {str(header or "").strip().lower(): str(header or "").strip() for header in headers}
    aliases = {
        "url_originale": ["url_originale", "url originale", "original_url"],
        "url_finale": ["url_finale", "url finale", "final_url"],
        "url": ["url"],
        "domaine": ["domaine", "domain"],
        "nom_source": ["nom_source", "source", "name"],
        "categorie": ["categorie", "catégorie", "categorie_principale"],
        "priorite": ["priorite", "priorité"],
    }
    result = {}
    for key, names in aliases.items():
        for name in names:
            if name in normalized:
                result[key] = normalized[name]
                break
    return result


def load_existing_catalog(path: Path) -> ExistingCatalog:
    wb = load_workbook(path, data_only=False)
    sheet_name = "Catalogue_liens" if "Catalogue_liens" in wb.sheetnames else wb.sheetnames[0]
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return ExistingCatalog([], set(), set(), set(), set(), set())
    headers = [str(value or "").strip() for value in rows[0]]
    hmap = _header_map(headers)
    records: list[dict] = []
    original_urls: set[str] = set()
    final_urls: set[str] = set()
    normalized_urls: set[str] = set()
    domains: set[str] = set()
    fingerprints: set[str] = set()
    for row in rows[1:]:
        record = {headers[i]: row[i] if i < len(row) else "" for i in range(len(headers))}
        urls = []
        for key in URL_COLUMNS:
            col = hmap.get(key)
            if col and record.get(col):
                urls.append(str(record[col]))
        if not urls:
            continue
        for url in urls:
            normalized = normalize_url(url)
            normalized_urls.add(normalized)
            fingerprints.add(url_fingerprint(normalized))
            domains.add(get_domain(normalized))
        if hmap.get("url_originale") and record.get(hmap["url_originale"]):
            original_urls.add(normalize_url(str(record[hmap["url_originale"]])))
        if hmap.get("url_finale") and record.get(hmap["url_finale"]):
            final_urls.add(normalize_url(str(record[hmap["url_finale"]])))
        record["Est_nouveau"] = "Non"
        records.append(record)
    return ExistingCatalog(records, original_urls, final_urls, normalized_urls, domains, fingerprints)
